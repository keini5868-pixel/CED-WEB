"""Inbound document ingest — PDF y Word (.docx) para chat / advanced."""

from __future__ import annotations

import logging
import re
import zipfile
from dataclasses import dataclass
from io import BytesIO

logger = logging.getLogger(__name__)

MAX_PDF_BYTES = 10 * 1024 * 1024
MAX_DOC_BYTES = MAX_PDF_BYTES
MAX_EXTRACT_CHARS = 50_000
MIN_USEFUL_CHARS = 20

# Alias histórico
MAX_DOCUMENT_BYTES = MAX_DOC_BYTES


class PdfIngestError(Exception):
    """Error de lectura / validación de documento entrante (PDF o Word)."""

    def __init__(self, message: str, *, http_status: int = 400):
        super().__init__(message)
        self.http_status = http_status


# Alias semántico
DocumentIngestError = PdfIngestError


@dataclass(frozen=True)
class PdfExtractResult:
    text: str
    filename: str
    page_count: int
    truncated: bool
    char_count: int
    kind: str = "pdf"  # "pdf" | "docx"


DocumentExtractResult = PdfExtractResult


def _ext(name: str | None) -> str:
    raw = (name or "").strip().lower()
    if "." not in raw:
        return ""
    return raw.rsplit(".", 1)[-1]


def _safe_filename(name: str | None, *, default_ext: str = "pdf") -> str:
    raw = (name or f"documento.{default_ext}").strip() or f"documento.{default_ext}"
    raw = raw.replace("\\", "/").split("/")[-1]
    raw = re.sub(r"[^\w.\- ()áéíóúÁÉÍÓÚñÑ]+", "_", raw, flags=re.UNICODE)
    ext = _ext(raw)
    if ext not in ("pdf", "docx"):
        raw = f"{raw}.{default_ext}"
    return raw[:120]


def detect_document_kind(data: bytes, filename: str | None = None) -> str:
    """Devuelve 'pdf' | 'docx' o lanza si no se reconoce."""
    name = (filename or "").lower()
    head = data[:8] if data else b""
    if head.startswith(b"%PDF") or name.endswith(".pdf"):
        if head and not head.startswith(b"%PDF") and name.endswith(".pdf"):
            # Extensión .pdf pero sin cabecera — aún intentamos PDF extract
            return "pdf"
        if head.startswith(b"%PDF"):
            return "pdf"
        if name.endswith(".pdf"):
            return "pdf"
    if name.endswith(".docx") or (
        head.startswith(b"PK") and _looks_like_docx_zip(data)
    ):
        return "docx"
    if name.endswith(".doc"):
        raise PdfIngestError(
            "Los .doc antiguos (Word 97-2003) no están soportados. "
            "Guarde el archivo como .docx o PDF e inténtelo de nuevo.",
        )
    raise PdfIngestError(
        "Formato no soportado. Sube un PDF o un Word (.docx).",
    )


def _looks_like_docx_zip(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            names = set(zf.namelist())
        return "word/document.xml" in names
    except Exception:  # noqa: BLE001
        return False


def validate_pdf_bytes(data: bytes, *, filename: str | None = None) -> None:
    if not data:
        raise PdfIngestError("El PDF está vacío.")
    if len(data) > MAX_PDF_BYTES:
        raise PdfIngestError(
            f"PDF demasiado grande. Máximo {MAX_PDF_BYTES // (1024 * 1024)} MB.",
        )
    head = data[:8]
    if not head.startswith(b"%PDF"):
        raise PdfIngestError(
            "El archivo no parece un PDF válido (falta la cabecera %PDF).",
        )
    _ = filename


def validate_docx_bytes(data: bytes, *, filename: str | None = None) -> None:
    if not data:
        raise PdfIngestError("El documento Word está vacío.")
    if len(data) > MAX_DOC_BYTES:
        raise PdfIngestError(
            f"Documento demasiado grande. Máximo {MAX_DOC_BYTES // (1024 * 1024)} MB.",
        )
    if not data.startswith(b"PK"):
        raise PdfIngestError(
            "El archivo no parece un .docx válido (debe ser Word moderno).",
        )
    if not _looks_like_docx_zip(data):
        raise PdfIngestError(
            "No pude leer ese .docx. Prueba a guardarlo de nuevo desde Word "
            "o exportarlo a PDF.",
        )
    _ = filename


def extract_pdf_text(
    data: bytes,
    *,
    filename: str | None = None,
    max_chars: int = MAX_EXTRACT_CHARS,
) -> PdfExtractResult:
    """Extrae texto seleccionable con pypdf. No OCR (escaneos → error claro)."""
    validate_pdf_bytes(data, filename=filename)
    safe_name = _safe_filename(filename, default_ext="pdf")

    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise PdfIngestError(
            "Servicio de lectura PDF no disponible (falta pypdf).",
            http_status=503,
        ) from exc

    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        logger.warning("[PDF-INGEST] open failed: %s", exc)
        raise PdfIngestError(
            "No pude abrir ese PDF. Prueba con otro archivo o exporta de nuevo a PDF.",
        ) from exc

    if getattr(reader, "is_encrypted", False):
        try:
            ok = reader.decrypt("")
            if ok == 0:
                raise PdfIngestError(
                    "Este PDF está protegido con contraseña. Quita la protección e inténtalo de nuevo.",
                )
        except PdfIngestError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise PdfIngestError(
                "Este PDF está protegido o no se puede desencriptar.",
            ) from exc

    pages = list(reader.pages or [])
    page_count = len(pages)
    if page_count == 0:
        raise PdfIngestError("El PDF no tiene páginas legibles.")

    parts: list[str] = []
    total = 0
    truncated = False
    for i, page in enumerate(pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            logger.warning("[PDF-INGEST] page %s extract failed name=%s", i, safe_name)
            raw = ""
        chunk = raw.strip()
        if not chunk:
            continue
        header = f"--- Página {i} ---\n"
        piece = f"{header}{chunk}"
        if total + len(piece) + 2 > max_chars:
            remain = max_chars - total - len(header) - 20
            if remain > 80:
                parts.append(f"{header}{chunk[:remain]}\n…[texto truncado]")
            truncated = True
            break
        parts.append(piece)
        total += len(piece) + 2

    text = "\n\n".join(parts).strip()
    body_only = re.sub(r"---\s*Página\s+\d+\s*---", "", text, flags=re.I)
    useful = re.sub(r"\s+", "", body_only)
    if len(useful) < MIN_USEFUL_CHARS:
        raise PdfIngestError(
            "Este PDF parece escaneado o sin texto seleccionable. "
            "Por ahora CED lee PDFs con texto copiable; si es escaneo, "
            "sube fotos claras de las páginas como imagen.",
        )

    return PdfExtractResult(
        text=text,
        filename=safe_name,
        page_count=page_count,
        truncated=truncated,
        char_count=len(text),
        kind="pdf",
    )


def extract_docx_text(
    data: bytes,
    *,
    filename: str | None = None,
    max_chars: int = MAX_EXTRACT_CHARS,
) -> PdfExtractResult:
    """Extrae párrafos y tablas de un .docx con python-docx."""
    validate_docx_bytes(data, filename=filename)
    safe_name = _safe_filename(filename, default_ext="docx")

    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise PdfIngestError(
            "Servicio de lectura Word no disponible (falta python-docx).",
            http_status=503,
        ) from exc

    try:
        doc = Document(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        logger.warning("[DOCX-INGEST] open failed: %s", exc)
        raise PdfIngestError(
            "No pude abrir ese Word (.docx). Guárdalo de nuevo o exporta a PDF.",
        ) from exc

    parts: list[str] = []
    for para in doc.paragraphs:
        chunk = (para.text or "").strip()
        if chunk:
            parts.append(chunk)

    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts).strip()
    useful = re.sub(r"\s+", "", text)
    if len(useful) < MIN_USEFUL_CHARS:
        raise PdfIngestError(
            "Este Word (.docx) no tiene texto legible. "
            "Si es solo imágenes, sube fotos claras o un PDF con texto seleccionable.",
        )

    truncated = False
    if len(text) > max_chars:
        text = text[: max_chars - 20].rstrip() + "\n…[texto truncado]"
        truncated = True

    # page_count ≈ bloques de ~1800 chars (heurística para el LLM)
    page_count = max(1, (len(text) + 1799) // 1800)

    return PdfExtractResult(
        text=text,
        filename=safe_name,
        page_count=page_count,
        truncated=truncated,
        char_count=len(text),
        kind="docx",
    )


def extract_document_text(
    data: bytes,
    *,
    filename: str | None = None,
    max_chars: int = MAX_EXTRACT_CHARS,
) -> PdfExtractResult:
    """Detecta PDF o .docx y extrae texto."""
    kind = detect_document_kind(data, filename)
    if kind == "docx":
        return extract_docx_text(data, filename=filename, max_chars=max_chars)
    return extract_pdf_text(data, filename=filename, max_chars=max_chars)


def format_pdf_for_llm(
    extracted: PdfExtractResult,
    user_message: str,
) -> str:
    """Bloque de contexto para el turno LLM (no se guarda entero en historial)."""
    note = ""
    if extracted.truncated:
        note = (
            "\n(Nota: el documento es largo; solo se incluye el inicio "
            f"hasta ~{MAX_EXTRACT_CHARS} caracteres.)\n"
        )
    kind = (extracted.kind or "pdf").lower()
    label = "WORD (.docx)" if kind == "docx" else "PDF"
    unit = "sección(es)" if kind == "docx" else "página(s)"
    caption = (user_message or "").strip() or (
        f"Analiza este documento {label}: resume lo importante, "
        "destaca puntos clave y responde con claridad."
    )
    return (
        f"[DOCUMENTO {label} ADJUNTO: {extracted.filename} — "
        f"{extracted.page_count} {unit}]\n"
        f"{note}"
        f"{extracted.text}\n"
        f"---\n"
        f"Pedido del usuario: {caption}"
    )


def user_display_for_pdf(filename: str, user_message: str) -> str:
    """Texto corto para historial / UI (sin volcar el documento entero)."""
    ext = _ext(filename)
    default_ext = "docx" if ext == "docx" else "pdf"
    name = _safe_filename(filename, default_ext=default_ext)
    icon = "📄"
    label = "Word" if name.lower().endswith(".docx") else "PDF"
    caption = (user_message or "").strip()
    if caption:
        return f"{icon} {label}: {name}\n{caption}"
    return f"{icon} {label}: {name}"
